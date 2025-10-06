"""Excel parser for catalog files."""

from typing import List

import openpyxl
import xlrd
from io import BytesIO


class ExcelParser:
    """Parser for Excel catalog files (.xlsx and .xls)."""

    # Column name mappings (same as CSV parser)
    TITLE_COLUMNS = ["title", "song", "work", "track"]
    ARTIST_COLUMNS = ["artist", "composer", "writer", "contributors"]
    DURATION_COLUMNS = ["duration", "length", "time"]
    ISWC_COLUMNS = ["iswc", "iswc_code"]

    def parse(self, file_content: bytes, filename: str) -> List[dict]:
        """Parse Excel file content into list of tracks.

        Args:
            file_content: Raw Excel file bytes
            filename: Original filename to determine format

        Returns:
            List of track dictionaries

        Raises:
            ValueError: If Excel is invalid or missing required columns
        """
        if filename.endswith(".xlsx"):
            return self._parse_xlsx(file_content)
        elif filename.endswith(".xls"):
            return self._parse_xls(file_content)
        else:
            raise ValueError(f"Unsupported Excel format: {filename}")

    def _parse_xlsx(self, file_content: bytes) -> List[dict]:
        """Parse .xlsx file using openpyxl.

        Args:
            file_content: Raw .xlsx file bytes

        Returns:
            List of track dictionaries
        """
        try:
            # Load workbook
            wb = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
            ws = wb.active

            # Get headers from first row
            headers = []
            for cell in ws[1]:
                headers.append(cell.value.lower().strip() if cell.value else "")

            # Find column mappings
            title_col = self._find_column(headers, self.TITLE_COLUMNS)
            artist_col = self._find_column(headers, self.ARTIST_COLUMNS)
            duration_col = self._find_column(headers, self.DURATION_COLUMNS)
            iswc_col = self._find_column(headers, self.ISWC_COLUMNS)

            if title_col is None:
                raise ValueError(
                    f"Excel must have a title column (one of: {', '.join(self.TITLE_COLUMNS)})"
                )

            title_idx = headers.index(title_col)
            artist_idx = headers.index(artist_col) if artist_col else None
            duration_idx = headers.index(duration_col) if duration_col else None
            iswc_idx = headers.index(iswc_col) if iswc_col else None

            # Parse rows
            tracks = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                # Get title
                title = row[title_idx] if title_idx < len(row) else None
                if not title or not str(title).strip():
                    continue

                track = {"title": str(title).strip()}

                # Get artist
                if artist_idx is not None and artist_idx < len(row):
                    artist = row[artist_idx]
                    if artist and str(artist).strip():
                        track["artist"] = str(artist).strip()

                # Get duration
                if duration_idx is not None and duration_idx < len(row):
                    duration = self._parse_duration(row[duration_idx])
                    if duration:
                        track["duration"] = duration

                # Get ISWC
                if iswc_idx is not None and iswc_idx < len(row):
                    iswc = row[iswc_idx]
                    if iswc and str(iswc).strip():
                        track["iswc"] = str(iswc).strip()

                tracks.append(track)

            if not tracks:
                raise ValueError("Excel file contains no valid track data")

            return tracks

        except Exception as e:
            raise ValueError(f"Error parsing .xlsx file: {str(e)}")

    def _parse_xls(self, file_content: bytes) -> List[dict]:
        """Parse .xls file using xlrd.

        Args:
            file_content: Raw .xls file bytes

        Returns:
            List of track dictionaries
        """
        try:
            # Load workbook
            wb = xlrd.open_workbook(file_contents=file_content)
            ws = wb.sheet_by_index(0)

            # Get headers from first row
            headers = []
            for col in range(ws.ncols):
                cell_value = ws.cell_value(0, col)
                headers.append(str(cell_value).lower().strip() if cell_value else "")

            # Find column mappings
            title_col = self._find_column(headers, self.TITLE_COLUMNS)
            artist_col = self._find_column(headers, self.ARTIST_COLUMNS)
            duration_col = self._find_column(headers, self.DURATION_COLUMNS)
            iswc_col = self._find_column(headers, self.ISWC_COLUMNS)

            if title_col is None:
                raise ValueError(
                    f"Excel must have a title column (one of: {', '.join(self.TITLE_COLUMNS)})"
                )

            title_idx = headers.index(title_col)
            artist_idx = headers.index(artist_col) if artist_col else None
            duration_idx = headers.index(duration_col) if duration_col else None
            iswc_idx = headers.index(iswc_col) if iswc_col else None

            # Parse rows
            tracks = []
            for row_num in range(1, ws.nrows):
                # Get title
                title = ws.cell_value(row_num, title_idx)
                if not title or not str(title).strip():
                    continue

                track = {"title": str(title).strip()}

                # Get artist
                if artist_idx is not None:
                    artist = ws.cell_value(row_num, artist_idx)
                    if artist and str(artist).strip():
                        track["artist"] = str(artist).strip()

                # Get duration
                if duration_idx is not None:
                    duration = self._parse_duration(ws.cell_value(row_num, duration_idx))
                    if duration:
                        track["duration"] = duration

                # Get ISWC
                if iswc_idx is not None:
                    iswc = ws.cell_value(row_num, iswc_idx)
                    if iswc and str(iswc).strip():
                        track["iswc"] = str(iswc).strip()

                tracks.append(track)

            if not tracks:
                raise ValueError("Excel file contains no valid track data")

            return tracks

        except Exception as e:
            raise ValueError(f"Error parsing .xls file: {str(e)}")

    def _find_column(self, headers: List[str], possible_names: List[str]) -> str | None:
        """Find matching column name from possible options."""
        for header in headers:
            if header in possible_names:
                return header
        return None

    def _parse_duration(self, value) -> int | None:
        """Parse duration value into seconds."""
        if value is None:
            return None

        # If already a number, return it
        if isinstance(value, (int, float)):
            return int(value)

        # If string, try parsing
        if isinstance(value, str):
            value = value.strip()
            try:
                return int(value)
            except ValueError:
                # Try parsing as time format
                parts = value.split(":")
                if len(parts) == 2:
                    minutes, seconds = parts
                    return int(minutes) * 60 + int(seconds)
                elif len(parts) == 3:
                    hours, minutes, seconds = parts
                    return int(hours) * 3600 + int(minutes) * 60 + int(seconds)

        return None
